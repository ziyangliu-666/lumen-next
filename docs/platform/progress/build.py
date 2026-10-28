"""Build the ASS fortnightly progress reports as Word files.

Each report has a Markdown source in this folder (2026-10-20.md, 2026-10-27.md).
This script writes "Team 15_<date>.docx" for each source. Run figures.py first:
it makes every figure in figures/.

The course template (00_Report-Template.docx) gives the styles, the page setup
and the cover look. The script keeps the template's cover paragraphs, fills
them, and removes the rest of the template body, because a progress report has
six sections of its own and no table of contents.

Run from any folder:
    python3 -I docs/platform/progress/figures.py
    python3 -I docs/platform/progress/build.py
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
TEMPLATE = (
    Path.home()
    / "Downloads/NUS-SE34-Coursework/ASS-Architecting-Scalable-Systems/00_Report-Template.docx"
)
MEMBERS = "HE RONG, LI ZHUOYANG, LIU ZIYANG, LU YI, PENG XUJIAN"
SOURCES = ["2026-10-20.md", "2026-10-27.md"]

# ---------------------------------------------------------------- Markdown


def parse(md: str) -> tuple[dict[str, str], list[tuple[str, object]]]:
    """Read the small Markdown subset the sources use.

    Front matter (title, date, period); "## " headings; paragraphs; pipe tables;
    a line "Table N. ..." as a table caption; "![caption](path)" as a figure.
    """
    meta: dict[str, str] = {}
    body = md
    fm = re.match(r"^---\n(.*?)\n---\n", md, re.S)
    if fm:
        for line in fm.group(1).splitlines():
            key, _, value = line.partition(":")
            meta[key.strip()] = value.strip()
        body = md[fm.end():]
    blocks: list[tuple[str, object]] = []
    for chunk in re.split(r"\n\s*\n", body.strip()):
        lines = chunk.strip().splitlines()
        # A caption line can sit directly above its table.
        if lines and lines[0].startswith("Table ") and len(lines) > 1 and lines[1].startswith("|"):
            caption = []
            while lines and not lines[0].startswith("|"):
                caption.append(lines.pop(0))
            blocks.append(("tcaption", " ".join(caption)))
        if not lines:
            continue
        first = lines[0]
        if first.startswith("## "):
            blocks.append(("h1", first[3:].strip()))
        elif first.startswith("!["):
            m = re.match(r"!\[(.*)\]\(([^)]*)\)(?:\{width=([\d.]+)\})?", " ".join(lines))
            blocks.append(("figure", (m.group(1), m.group(2), float(m.group(3)) if m.group(3) else None)))
        elif first.startswith("|"):
            rows = [
                [c.strip() for c in ln.strip().strip("|").split("|")]
                for ln in lines
                if not re.match(r"^\|[-| :]+\|$", ln.strip())
            ]
            blocks.append(("table", rows))
        elif first.startswith("Table ") and len(lines) >= 1 and all(not ln.startswith("|") for ln in lines):
            blocks.append(("tcaption", " ".join(lines)))
        else:
            blocks.append(("p", " ".join(ln.strip() for ln in lines)))
    return meta, blocks


def add_runs(par, text: str, size: float | None = None) -> None:
    """Add text with **bold** for a defined term and `code` in a mono font."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            run = par.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Courier New"
        else:
            run = par.add_run(part)
        if size:
            run.font.size = Pt(size)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(table) -> None:
    """The template's own table look: outer 1 pt, inner 0.5 pt, black."""
    tbl_pr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge, size in (("top", 8), ("left", 8), ("bottom", 8), ("right", 8), ("insideH", 4), ("insideV", 4)):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(size))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "auto")
        b.append(e)
    tbl_pr.append(b)


def fixed_grid(table, widths_cm: list[float]) -> None:
    """Write the column widths into the table grid and fix the layout.

    Cell widths alone are ignored by some renderers (LibreOffice), which then
    share the width equally and wrap short columns badly.
    """
    twips = [int(w * 567) for w in widths_cm]
    tbl_pr = table._tbl.tblPr
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(twips)))
    tbl_w.set(qn("w:type"), "dxa")
    for old in tbl_pr.findall(qn("w:tblW")):
        tbl_pr.remove(old)
    tbl_pr.append(tbl_w)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    grid = table._tbl.tblGrid
    for col, w in zip(grid.findall(qn("w:gridCol")), twips):
        col.set(qn("w:w"), str(w))


def cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:cantSplit")
    e.set(qn("w:val"), "true")
    tr_pr.append(e)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    tr_pr.append(h)


def caption(doc, text: str, before: float = 4, after: float = 4):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(before)
    par.paragraph_format.space_after = Pt(after)
    par.paragraph_format.keep_with_next = text.startswith("Table")
    m = re.match(r"((?:Figure|Table) \d+\.)(.*)", text)
    label, rest = (m.group(1), m.group(2)) if m else ("", text)
    r = par.add_run(label)
    r.bold = True
    r.font.size = Pt(9.5)
    add_runs(par, rest, size=9.5)
    return par


def column_widths(rows: list[list[str]], total_cm: float) -> list[float]:
    """Share the text width by the longest text in each column, with a floor."""
    cols = range(len(rows[0]))
    longest = [max(len(r[i]) for r in rows) for i in cols]
    # A column is never narrower than its longest word, so names do not break.
    word = [max(len(w) for r in rows for w in r[i].split(" ")) for i in cols]
    weights = [max(min(n, 60), w * 2.6) for n, w in zip(longest, word)]
    s = sum(weights)
    return [total_cm * w / s for w in weights]


def build(source: Path) -> Path:
    meta, blocks = parse(source.read_text())
    doc = Document(str(TEMPLATE))
    body = doc.element.body

    # Keep the cover (template paragraphs 0 to 13), drop the rest of the template body.
    paras = doc.paragraphs
    keep = {id(p._p) for p in paras[:14]}
    for child in list(body):
        if child.tag == qn("w:sectPr") or id(child) in keep:
            continue
        body.remove(child)

    def set_text(p, text: str) -> None:
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        if p.runs:
            p.runs[0].text = text
        else:
            p.add_run(text)

    paras = doc.paragraphs
    set_text(paras[4], meta["title"])
    set_text(paras[6], "Practice Module for Graduate Certificate in Architecting Scalable Systems")
    set_text(paras[9], "Team 15")
    set_text(paras[10], "Members:")
    set_text(paras[11], MEMBERS)
    set_text(paras[12], f"Date: {meta['date']}")
    paras[12].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_text(paras[13], f"Period: {meta['period']}")
    paras[13].alignment = WD_ALIGN_PARAGRAPH.CENTER
    paras[13].add_run().add_break(WD_BREAK.PAGE)

    section = doc.sections[0]
    text_cm = (section.page_width - section.left_margin - section.right_margin) / 360000

    normal = doc.styles["Normal"]
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for kind, value in blocks:
        if kind == "h1":
            doc.add_paragraph(value, style="Heading 1").paragraph_format.space_after = Pt(6)
        elif kind == "p":
            add_runs(doc.add_paragraph(), value, size=11)
        elif kind == "tcaption":
            caption(doc, value, before=8)
        elif kind == "figure":
            cap, rel, width_cm = value
            path = (HERE / rel).resolve()
            pic = doc.add_paragraph()
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.paragraph_format.keep_with_next = True
            pic.paragraph_format.space_before = Pt(6)
            pic.add_run().add_picture(str(path), width=Cm(width_cm or text_cm))
            caption(doc, cap, after=8)
        elif kind == "table":
            # A short cell (a date, a name) stays on one line.
            rows = [[c.replace(" ", "\u00a0") if len(c) <= 12 else c for c in r] for r in value]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            borders(table)
            table.autofit = False
            widths = column_widths(rows, text_cm)
            fixed_grid(table, widths)
            for i, row in enumerate(rows):
                cant_split(table.rows[i])
                last_total = i == len(rows) - 1 and row[0] == "Total"
                for j, text in enumerate(row):
                    cell = table.cell(i, j)
                    cell.width = Cm(widths[j])
                    par = cell.paragraphs[0]
                    par.paragraph_format.space_after = Pt(0)
                    # Keep a short table on one page: every row holds to the next.
                    par.paragraph_format.keep_with_next = i < len(rows) - 1 and len(rows) <= 12
                    par.paragraph_format.line_spacing = 1.0
                    if re.fullmatch(r"\d[\d,]*", text):
                        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    add_runs(par, text, size=9.5)
                    if i == 0 or last_total:
                        for r in par.runs:
                            r.bold = True
                    if i == 0:
                        shade(cell, "D9D9D9")
            repeat_header(table.rows[0])
            if (kind, value) is not blocks[-1]:
                spacer = doc.add_paragraph()
                spacer.paragraph_format.space_after = Pt(2)

    out = HERE / f"Team 15_{source.stem}.docx"
    doc.save(str(out))
    return out


def main() -> None:
    for name in SOURCES:
        print(build(HERE / name))


if __name__ == "__main__":
    main()
