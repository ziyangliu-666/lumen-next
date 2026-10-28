"""Builds the Word report from report.md and the course template.

    python3 docs/platform/report/build_docx.py

Input:
- report.md (this folder): the report source. "##" is a Heading 1, "###" a
  Heading 2 and "####" a Heading 3; the template's heading styles number them.
- the course template, 00_Report-Template.docx: its cover page, styles,
  footer and table of contents field are kept; its instruction text and
  example tables are removed.
- figures/*.png: the diagrams and screenshots that report.md names.
  A figure line ending in {wide} goes on a landscape page.

Output: out/Team 15 - Lumen Platform Report.docx and, when LibreOffice is
installed, the PDF next to it. The page numbers in the table of contents are
read from a first PDF render and written in a second pass, so the PDF shows
them; Word refreshes the field when the file is opened.
"""
import copy
import os
import re
import shutil
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.expanduser(
    os.environ.get("REPORT_TEMPLATE",
                   "~/Downloads/NUS-SE34-Coursework/ASS-Architecting-Scalable-Systems/00_Report-Template.docx"))
OUT_DIR = os.path.join(HERE, "out")
NAME = "Team 15 - Lumen Platform Report"
SOFFICE = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"

TITLE = "Project Report for Lumen Platform"
TEAM = "Team 15"
MEMBERS = ["HE RONG", "LI ZHUOYANG", "LIU ZIYANG", "LU YI", "PENG XUJIAN"]
EXTRA_COVER = ["Public address: https://lumen.139-180-185-51.sslip.io",
               "Code base: https://github.com/ziyangliu-666/lumen-next (branch ass/platform)",
               "November 2026"]

TEXT_W = Inches(6.27)            # A4 portrait, 1 inch margins
LAND_W = Inches(9.69)            # A4 landscape, 1 inch margins
LAND_H = Inches(5.6)             # height left for figures and captions on a landscape page
GREY = RGBColor(0x4F, 0x5D, 0x75)
CODE_FONT = "Courier New"


# ---------------------------------------------------------------- small XML helpers
def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge, size in (("top", 8), ("left", 8), ("bottom", 8), ("right", 8), ("insideH", 4), ("insideV", 4)):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tblPr.append(borders)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def field_run(paragraph, instr):
    """Appends a complete field (begin, instruction, separate, end) to a paragraph."""
    for kind, text in (("begin", None), (None, instr), ("separate", None), ("end", None)):
        r = paragraph.add_run()
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = instr
            r._r.append(it)


# ---------------------------------------------------------------- inline Markdown
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")


def add_inline(paragraph, text, size=None, color=None, bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = CODE_FONT
            run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), CODE_FONT)
            if size:
                run.font.size = Pt(size.pt - 1 if hasattr(size, "pt") else size - 1)
            else:
                run.font.size = Pt(10.5)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if size and not part.startswith("`"):
            run.font.size = size if hasattr(size, "pt") else Pt(size)
        if color:
            run.font.color.rgb = color
        if bold:
            run.bold = True


# ---------------------------------------------------------------- Markdown blocks
def parse(md):
    """Yields (kind, data) blocks from the report source."""
    lines = md.split("\n")
    i = 0
    # the cover lines before the first section are written from the constants above
    while i < len(lines) and not lines[i].startswith("## "):
        i += 1
    while i < len(lines):
        line = lines[i]
        if not line.strip() or line.strip() == "---":
            i += 1
            continue
        m = re.match(r"^(#{2,4}) (.*)$", line)
        if m:
            level = len(m.group(1)) - 1
            text = re.sub(r"^\d+(\.\d+)*\.?\s+", "", m.group(2))
            yield ("heading", (level, text))
            i += 1
            continue
        m = re.match(r"^!\[(.*)\]\((.*?)\)(\{wide\})?\s*$", line)
        if m:
            yield ("figure", (m.group(1), m.group(2), bool(m.group(3))))
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.match(r"^:?-+:?$", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            yield ("table", rows)
            continue
        m = re.match(r"^(\s*)([-*]|\d+\.) (.*)$", line)
        if m:
            items = []
            while i < len(lines):
                m = re.match(r"^(\s*)([-*]|\d+\.) (.*)$", lines[i])
                if m:
                    items.append([len(m.group(1)) // 2, m.group(2), m.group(3)])
                    i += 1
                elif lines[i].startswith("  ") and lines[i].strip() and items:
                    items[-1][2] += " " + lines[i].strip()
                    i += 1
                else:
                    break
            yield ("list", items)
            continue
        para = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{2,4} |!\[|\||\s*([-*]|\d+\.) |---)", lines[i]):
            para.append(lines[i].strip())
            i += 1
        yield ("para", " ".join(para))


# ---------------------------------------------------------------- the document
class Builder:
    def __init__(self, toc_pages=None):
        self.doc = Document(TEMPLATE)
        self.toc_pages = toc_pages or {}
        self.headings = []          # (level, number, text)
        self.counters = [0, 0, 0]
        self.body = self.doc.element.body
        self._prepare()

    # -------------------------------------------------------- template
    def _prepare(self):
        d = self.doc
        paras = d.paragraphs
        # cover
        self._set_text(paras[4], TITLE)
        self._set_text(paras[9], TEAM)
        self._set_text(paras[11], ", ".join(MEMBERS))
        for i, line in enumerate(EXTRA_COVER):
            p = paras[12 + i] if 12 + i < 14 else None
            if p is not None:
                self._set_text(p, line)
        # the TOC heading stays; everything from the first TOC entry onward goes
        contents = paras[15]
        self.toc_anchor = contents
        sect = self.body.find(qn("w:sectPr"))
        started = False
        for el in list(self.body):
            if el is contents._p:
                started = True
                continue
            if started and el is not sect:
                self.body.remove(el)
        # styles that the template does not have
        styles = d.styles
        if "Caption" not in [s.name for s in styles]:
            cap = styles.add_style("Caption", 1)
        else:
            cap = styles["Caption"]
        cap.base_style = styles["Normal"]
        cap.font.size = Pt(9.5)
        cap.font.italic = False
        cap.font.color.rgb = GREY
        cap.paragraph_format.space_before = Pt(4)
        cap.paragraph_format.space_after = Pt(14)
        normal = styles["Normal"]
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.08
        # Word refreshes the TOC field when the file is opened
        settings = d.settings.element
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        settings.append(upd)

    @staticmethod
    def _set_text(p, text):
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        if p.runs:
            p.runs[0].text = text
        else:
            p.add_run(text)

    # -------------------------------------------------------- blocks
    def heading(self, level, text):
        self.counters[level - 1] += 1
        for k in range(level, 3):
            self.counters[k] = 0
        number = ".".join(str(c) for c in self.counters[:level])
        self.headings.append((level, number, text))
        # no forced page break before a chapter: it left half-empty pages at the end of
        # sections 2 to 5; keep_with_next keeps a heading off the bottom of a page
        p = self.doc.add_paragraph(style="Heading %d" % level)
        p.add_run(text)
        bm = "_Toc%04d" % len(self.headings)
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(len(self.headings)))
        start.set(qn("w:name"), bm)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(len(self.headings)))
        p._p.insert(1, start)
        p._p.append(end)
        p.paragraph_format.space_before = Pt(22 if level == 1 else 10)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True

    def para(self, text):
        p = self.doc.add_paragraph()
        add_inline(p, text)
        if text.rstrip().endswith(":") or text.startswith("**") and text.rstrip().endswith("**"):
            p.paragraph_format.keep_with_next = True

    def bullet_list(self, items):
        for k, (depth, mark, text) in enumerate(items):
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            hang = 0.22 if mark in "-*" else 0.32      # room for "10." and "12."
            indent = 0.1 + hang + 0.28 * depth
            pf.left_indent = Inches(indent)
            pf.first_line_indent = Inches(-hang)
            pf.space_after = Pt(8 if k == len(items) - 1 else 3)
            pf.tab_stops.add_tab_stop(Inches(indent))
            p.add_run(("•" if mark in "-*" else mark) + "\t")
            add_inline(p, text)

    def table(self, rows):
        ncols = max(len(r) for r in rows)
        rows = [r + [""] * (ncols - len(r)) for r in rows]
        t = self.doc.add_table(rows=len(rows), cols=ncols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t)
        # column widths: each column gets at least its longest word (a short label column
        # gets its whole text, so "reader agent" or "8 to 19 Oct" does not wrap), and the
        # rest of the width goes to the columns with the most text
        size = Pt(9) if ncols >= 4 or sum(len(c) for r in rows for c in r) > 1500 else Pt(9.5)
        plain = [[re.sub(r"[*`]", "", c) for c in r] for r in rows]
        pad = Inches(0.16)                      # left and right cell margins together
        avail = TEXT_W - pad * ncols
        total_chars = avail / 12700 / (size.pt * 0.55)
        def wlen(text, bold=False):
            # capitals and bold text are wider than the average character
            n = len(text) + 0.5 * sum(ch.isupper() for ch in text)
            return n * 1.1 if bold else n
        mins = []
        for c in range(ncols):
            longest_word = max(max((wlen(w, ri == 0) for w in r[c].split()), default=3) for ri, r in enumerate(plain))
            longest_cell = max(wlen(r[c], ri == 0) for ri, r in enumerate(plain))
            mins.append(min(longest_cell if longest_cell <= 16 else longest_word, 24) + 1.5)
        prefs = [min(max(len(r[c]) for r in plain), 120) for c in range(ncols)]
        if sum(mins) >= total_chars:
            lens = mins
        else:
            extra = total_chars - sum(mins)
            want = [max(p - m, 0) for p, m in zip(prefs, mins)]
            tw = sum(want) or 1
            lens = [m + extra * w / tw for m, w in zip(mins, want)]
        if ncols == 2 and rows[0][0] == "Identifier":
            lens = [12, 88]
        if ncols == 6 and rows[0][0] == "Identifier":      # the AISS table
            lens = [10, 21, 23, 23, 13, 10]
        total = sum(lens)
        widths = [int(avail * n / total + pad) for n in lens]
        t.autofit = False
        tblPr = t._tbl.tblPr
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW")
            tblPr.append(tblW)
        tblW.set(qn("w:w"), str(int(TEXT_W / 635)))
        tblW.set(qn("w:type"), "dxa")
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w / 635)))
        short = len(rows) <= 8 and sum(len(c) for r in plain for c in r) <= 700
        for ri, row in enumerate(rows):
            tr = t.rows[ri]
            if ri == 0:
                repeat_header(tr)
            if len(" ".join(row)) < 1500:      # a row breaks across pages only when it is very long
                no_split(tr)
            for ci, text in enumerate(row):
                cell = tr.cells[ci]
                cell.width = widths[ci]
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                add_inline(p, text, size=size, bold=(ri == 0))
                # a short table stays on one page; a longer one may break, so that it does not
                # leave a large empty space; its header stays with the first row
                if ri == 0 or (short and ri < len(rows) - 1):
                    p.paragraph_format.keep_with_next = True
                if ri == 0:
                    set_cell_shading(cell, "F2F2F2")
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)

    def figure(self, caption, path, wide):
        from PIL import Image
        img = os.path.join(HERE, path)
        with Image.open(img) as im:
            w, h = im.size
        if wide:
            self._landscape(True)
            width = LAND_W
            if width * h / w > LAND_H:
                width = int(LAND_H * w / h)
        else:
            width = TEXT_W
            max_h = Inches(8.3)
            if width * h / w > max_h:
                width = int(max_h * w / h)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.add_run().add_picture(img, width=Emu(int(width)))
        c = self.doc.add_paragraph(style="Caption")
        m = re.match(r"^(Figure \d+\.)\s*(.*)$", caption)
        if m:
            r = c.add_run(m.group(1) + " ")
            r.bold = True
            add_inline(c, m.group(2))
        else:
            add_inline(c, caption)
        if wide:
            self._landscape(False)

    def _landscape(self, on):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        if on:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Inches(11.69), Inches(8.27)
        else:
            sec.orientation = WD_ORIENT.PORTRAIT
            sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Inches(1))

    # -------------------------------------------------------- TOC
    def toc(self):
        """Writes the TOC entries after the CONTENTS line, inside one TOC field."""
        anchor = self.toc_anchor._p
        style_ids = {1: "TOC1", 2: "TOC2", 3: "TOC3"}
        made = []
        for idx, (level, number, text) in enumerate(self.headings):
            p = OxmlElement("w:p")
            pPr = OxmlElement("w:pPr")
            ps = OxmlElement("w:pStyle")
            ps.set(qn("w:val"), style_ids[level])
            pPr.append(ps)
            tabs = OxmlElement("w:tabs")
            for kind, pos, leader in (("left", {1: 425, 2: 992, 3: 1560}[level], None), ("right", 9020, "dot")):
                tab = OxmlElement("w:tab")
                tab.set(qn("w:val"), kind)
                tab.set(qn("w:pos"), str(pos))
                if leader:
                    tab.set(qn("w:leader"), leader)
                tabs.append(tab)
            pPr.append(tabs)
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str({1: 0, 2: 284, 3: 568}[level]))
            pPr.append(ind)
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:after"), "60")
            pPr.append(sp)
            p.append(pPr)
            made.append(p)
        for p in reversed(made):
            anchor.addnext(p)
        from docx.text.paragraph import Paragraph
        paras = [Paragraph(p, self.toc_anchor._parent) for p in made]
        for idx, ((level, number, text), para) in enumerate(zip(self.headings, paras)):
            if idx == 0:
                r = para.add_run()
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); r._r.append(fc)
                r = para.add_run()
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = ' TOC \\o "1-3" \\h \\z \\u '; r._r.append(it)
                r = para.add_run()
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); r._r.append(fc)
            label = number + ("." if level == 1 else "")
            para.add_run(label + "\t" + text + "\t" + str(self.toc_pages.get(idx, "")))
            if idx == len(self.headings) - 1:
                r = para.add_run()
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r._r.append(fc)
        # the body starts on a new page after the contents
        last = paras[-1]._p
        brk = OxmlElement("w:p")
        r = OxmlElement("w:r")
        b = OxmlElement("w:br")
        b.set(qn("w:type"), "page")
        r.append(b)
        brk.append(r)
        last.addnext(brk)

    def build(self, md):
        first_h1 = True
        for kind, data in parse(md):
            if kind == "heading":
                level, text = data
                if level == 1 and first_h1:
                    first_h1 = False
                    self.counters = [0, 0, 0]
                self.heading(level, text)
            elif kind == "para":
                self.para(data)
            elif kind == "list":
                self.bullet_list(data)
            elif kind == "table":
                self.table(data)
            elif kind == "figure":
                self.figure(*data)
        self.toc()
        # the template's first section break carries the footer; keep A4 portrait
        return self.doc


# ---------------------------------------------------------------- page numbers from a PDF render
def to_pdf(docx_path, out_dir):
    profile = "file://" + tempfile.mkdtemp(prefix="lo_profile_ass_report_")
    subprocess.run([SOFFICE, "-env:UserInstallation=" + profile, "--headless", "--convert-to", "pdf",
                    "--outdir", out_dir, docx_path], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return os.path.join(out_dir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")


def heading_pages(pdf, headings):
    n = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", pdf], capture_output=True, text=True).stdout).group(1))
    pages = []
    for k in range(1, n + 1):
        txt = subprocess.run(["pdftotext", "-f", str(k), "-l", str(k), "-layout", pdf, "-"],
                             capture_output=True, text=True).stdout
        pages.append([re.sub(r"\s+", " ", l).strip() for l in txt.split("\n")])
    # skip the cover and the contents pages: start after the last page that holds "CONTENTS"
    start = max(i for i, pg in enumerate(pages) if any(l == "CONTENTS" for l in pg)) + 1
    while start < n and sum(1 for l in pages[start] if re.search(r"\.{5,}\s*\d*$", l)) > 3:
        start += 1
    found = {}
    cur = start
    for idx, (level, number, text) in enumerate(headings):
        pat = re.compile(r"^%s\.?\s+%s$" % (re.escape(number), re.escape(text)))
        for p in range(cur, n):
            if any(pat.match(l) for l in pages[p]):
                found[idx] = p + 1
                cur = p
                break
    return found


def main():
    with open(os.path.join(HERE, "report.md")) as f:
        md = f.read()
    os.makedirs(OUT_DIR, exist_ok=True)
    path = os.path.join(OUT_DIR, NAME + ".docx")
    b = Builder()
    b.build(md).save(path)
    if not os.path.exists(SOFFICE) or "--no-pdf" in sys.argv:
        print("wrote", path, "(no PDF)")
        return
    pdf = to_pdf(path, OUT_DIR)
    pages = heading_pages(pdf, b.headings)
    missing = [b.headings[i] for i in range(len(b.headings)) if i not in pages]
    b2 = Builder(pages)
    b2.build(md).save(path)
    pdf = to_pdf(path, OUT_DIR)
    print("wrote", path)
    print("wrote", pdf)
    if missing:
        print("headings not found in the PDF:", missing)


if __name__ == "__main__":
    main()
